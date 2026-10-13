"""Konversi lanjutan ala iLovePDF — 100% offline, tanpa internet.

Menyediakan:
  - PDF -> Word (.docx)      : teks + gambar, via python-docx
  - PDF -> PowerPoint (.pptx): tiap halaman jadi slide, via python-pptx
  - PDF -> Excel (.xlsx)     : tabel (find_tables) atau teks, via openpyxl
  - PDF -> PDF/A             : via Ghostscript bila ada, fallback best-effort
  - Ambil gambar dari PDF    : ekstrak gambar tertanam
  - Banding PDF (Compare)    : halaman berdampingan + sorot beda
  - Formulir (AcroForm)      : daftar + isi field
  - Scan -> PDF              : gambar -> PDF dengan perbaikan kontras

Dependensi berat diimpor lazy; bila absen -> ConvertError pesan ramah.
Semua fungsi headless-safe (tanpa tkinter) dan aman dari thread worker.
"""
from __future__ import annotations

import io
import os

import pymupdf
from PIL import Image

try:  # via app (src/ di path)
    from src.core.pdf_ops import _ensure_parent, _require_exists  # type: ignore
except ImportError:  # via tests (src/ di path)
    from core.pdf_ops import _ensure_parent, _require_exists  # type: ignore


class ConvertError(RuntimeError):
    pass


def _need_dep(mod: str, pip_name: str):
    try:
        return __import__(mod)
    except ImportError as e:  # noqa: BLE001
        raise ConvertError(
            f"Butuh paket '{pip_name}'. Pasang: pip install {pip_name}") from e


def _root() -> str:
    import sys
    here = os.path.dirname(os.path.abspath(__file__))
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.abspath(os.path.join(here, "..", ".."))


# ---------------------------------------------------------------- PDF -> Word
def pdf_to_word(pdf_path: str, output: str) -> str:
    """PDF -> DOCX: teks per blok + gambar tertanam (best-effort layout)."""
    _require_exists(pdf_path)
    _ensure_parent(output)
    docx = _need_dep("docx", "python-docx")
    from docx.shared import Inches

    document = docx.Document()
    with pymupdf.open(pdf_path) as src:
        if src.is_encrypted:
            raise ConvertError("PDF terkunci — buka dulu lewat Unlock.")
        for pno in range(src.page_count):
            page = src[pno]
            for blk in page.get_text("dict").get("blocks", []):
                if blk.get("type") == 0:
                    for line in blk.get("lines", []):
                        text = "".join(s.get("text", "") for s in line.get("spans", []))
                        if text.strip():
                            document.add_paragraph(text)
                elif blk.get("type") == 1 and blk.get("image"):
                    try:
                        document.add_picture(io.BytesIO(blk["image"]), width=Inches(5.5))
                    except Exception:  # noqa: BLE001
                        continue
            if pno < src.page_count - 1:
                document.add_page_break()
    document.save(output)
    return output


# ---------------------------------------------------------------- PDF -> PowerPoint
def pdf_to_powerpoint(pdf_path: str, output: str, dpi: int = 150) -> str:
    """PDF -> PPTX: tiap halaman jadi satu slide (gambar penuh)."""
    _require_exists(pdf_path)
    _ensure_parent(output)
    pptx = _need_dep("pptx", "python-pptx")
    from pptx.util import Emu

    if not 36 <= int(dpi) <= 300:
        raise ConvertError("DPI 36-300.")
    zoom = int(dpi) / 72.0
    with pymupdf.open(pdf_path) as src:
        if src.is_encrypted:
            raise ConvertError("PDF terkunci — buka dulu lewat Unlock.")
        if src.page_count == 0:
            raise ConvertError("PDF kosong.")
        prs = pptx.Presentation()
        first = src[0].rect
        prs.slide_width = Emu(int(first.width * 12700))   # 1pt = 12700 EMU
        prs.slide_height = Emu(int(first.height * 12700))
        blank = prs.slide_layouts[6]
        for page in src:
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
            buf = io.BytesIO(pix.tobytes("png"))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(buf, 0, 0,
                                     width=prs.slide_width, height=prs.slide_height)
        prs.save(output)
    return output


# ---------------------------------------------------------------- PDF -> Excel
def pdf_to_excel(pdf_path: str, output: str) -> dict:
    """PDF -> XLSX. Tiap tabel jadi sheet; halaman tanpa tabel -> baris teks."""
    _require_exists(pdf_path)
    _ensure_parent(output)
    openpyxl = _need_dep("openpyxl", "openpyxl")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n_tables = 0
    with pymupdf.open(pdf_path) as src:
        if src.is_encrypted:
            raise ConvertError("PDF terkunci — buka dulu lewat Unlock.")
        for pno in range(src.page_count):
            page = src[pno]
            try:
                tables = list(getattr(page.find_tables(), "tables", []) or [])
            except Exception:  # noqa: BLE001
                tables = []
            if tables:
                for ti, tbl in enumerate(tables):
                    rows = tbl.extract() or []
                    if not rows:
                        continue
                    ws = wb.create_sheet(title=f"Hal{pno + 1}-{ti + 1}"[:31])
                    for row in rows:
                        ws.append(["" if c is None else str(c) for c in row])
                    n_tables += 1
            else:
                lines = [ln for ln in page.get_text("text").splitlines() if ln.strip()]
                if lines:
                    ws = wb.create_sheet(title=f"Hal{pno + 1}"[:31])
                    for ln in lines:
                        ws.append([ln])
                    n_tables += 1
    if not wb.sheetnames:
        wb.create_sheet("kosong")
    wb.save(output)
    return {"output": output, "sheets": len(wb.sheetnames)}


# ---------------------------------------------------------------- PDF -> PDF/A
def find_ghostscript() -> str | None:
    import shutil
    cands = [
        os.path.join(_root(), "plugins", "ghostscript", "bin", "gswin64c.exe"),
        os.path.join(_root(), "plugins", "ghostscript", "bin", "gswin32c.exe"),
        os.path.join(_root(), "plugins", "ghostscript", "gswin64c.exe"),
        os.path.join(_root(), "plugins", "ghostscript", "gs"),
    ]
    for c in cands:
        if os.path.isfile(c):
            return c
    return shutil.which("gswin64c") or shutil.which("gswin32c") or shutil.which("gs")


def to_pdfa(pdf_path: str, output: str, level: int = 2) -> dict:
    """PDF -> PDF/A. Pakai Ghostscript bila ada (valid), else best-effort XMP."""
    _require_exists(pdf_path)
    _ensure_parent(output)
    if level not in (1, 2, 3):
        level = 2
    gs = find_ghostscript()
    if gs:
        import subprocess
        cmd = [
            gs, "-q", "-dNOPAUSE", "-dBATCH", "-dSAFER",
            "-sDEVICE=pdfwrite",
            f"-dPDFA={level}", "-dPDFACompatibilityPolicy=1",
            "-sColorConversionStrategy=RGB",
            f"-sOutputFile={output}", pdf_path,
        ]
        cp = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
        if cp.returncode == 0 and os.path.isfile(output):
            return {"output": output, "method": "ghostscript", "valid": True}
        # gs gagal -> jatuh ke best-effort
    _pdfa_best_effort(pdf_path, output, level)
    return {"output": output, "method": "best-effort", "valid": False}


def _pdfa_best_effort(pdf_path: str, output: str, level: int) -> None:
    xmp = (
        '<?xpacket begin="\ufeff" id="W5M0MpCehiHzreSzNTczkc9d"?>'
        '<x:xmpmeta xmlns:x="adobe:ns:meta/">'
        '<rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">'
        '<rdf:Description rdf:about="" '
        'xmlns:pdfaid="http://www.aiim.org/pdfa/ns/id/">'
        f"<pdfaid:part>{level}</pdfaid:part>"
        "<pdfaid:conformance>B</pdfaid:conformance>"
        "</rdf:Description></rdf:RDF></x:xmpmeta>"
        '<?xpacket end="w"?>'
    )
    with pymupdf.open(pdf_path) as doc:
        if doc.is_encrypted:
            raise ConvertError("PDF terkunci — buka dulu lewat Unlock.")
        try:
            doc.set_xml_metadata(xmp)
        except Exception:  # noqa: BLE001
            pass
        doc.save(output, garbage=4, deflate=True, clean=True)


# ---------------------------------------------------------------- extract images
def extract_images(pdf_path: str, out_dir: str) -> list[str]:
    """Ambil semua gambar tertanam dari PDF -> folder. Return daftar file."""
    _require_exists(pdf_path)
    os.makedirs(out_dir, exist_ok=True)
    made: list[str] = []
    with pymupdf.open(pdf_path) as doc:
        if doc.is_encrypted:
            raise ConvertError("PDF terkunci — buka dulu lewat Unlock.")
        for pno, page in enumerate(doc):
            seen: set[int] = set()
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                try:
                    d = doc.extract_image(xref)
                except Exception:  # noqa: BLE001
                    continue
                ext = (d.get("ext") or "png").lower()
                fp = os.path.join(out_dir, f"gambar-hal{pno + 1:03d}-{xref}.{ext}")
                with open(fp, "wb") as f:
                    f.write(d["image"])
                made.append(fp)
    if not made:
        raise ConvertError("Tidak ada gambar tertanam di PDF ini.")
    return made


# ---------------------------------------------------------------- compare
def compare_pdfs(a: str, b: str, output: str, dpi: int = 100) -> dict:
    """Banding halaman berdampingan + sorot area beda. Return ringkasan."""
    _require_exists(a)
    _require_exists(b)
    _ensure_parent(output)
    if not 50 <= int(dpi) <= 200:
        raise ConvertError("DPI 50-200.")
    from PIL import ImageChops, ImageDraw
    zoom = int(dpi) / 72.0
    GAP, BAND = 16, 26
    changed_pages: list[int] = []

    def _render(doc, i):
        if i >= doc.page_count:
            return None
        pix = doc[i].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
        return Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")

    with pymupdf.open(a) as da, pymupdf.open(b) as db:
        for d in (da, db):
            if d.is_encrypted:
                raise ConvertError("Ada PDF terkunci — buka dulu lewat Unlock.")
        total = max(da.page_count, db.page_count)
        if total == 0:
            raise ConvertError("PDF kosong.")
        out = pymupdf.open()
        for i in range(total):
            ia, ib = _render(da, i), _render(db, i)
            if ia is None:
                ia = Image.new("RGB", ib.size, (235, 235, 235))
            if ib is None:
                ib = Image.new("RGB", ia.size, (235, 235, 235))
            w = max(ia.width, ib.width)
            h = max(ia.height, ib.height)
            if (ia.width, ia.height) != (w, h):
                ia = ia.resize((w, h))
            if (ib.width, ib.height) != (w, h):
                ib = ib.resize((w, h))
            diff = ImageChops.difference(ia, ib).convert("L")
            ratio = 0.0
            boxes = []
            try:
                mask = diff.point(lambda v: 255 if v > 40 else 0)
                bbox = mask.getbbox()
                if bbox:
                    hist = mask.histogram()
                    nz = hist[255] if len(hist) > 255 else 0
                    ratio = round(100.0 * nz / (w * h), 2)
                    # kotak-kotak perubahan diskrit
                    gx, gy = 8, 10
                    cw, ch = w / gx, h / gy
                    for cx in range(gx):
                        for cy in range(gy):
                            cell = mask.crop((int(cx * cw), int(cy * ch),
                                              int((cx + 1) * cw), int((cy + 1) * ch)))
                            if cell.getbbox():
                                boxes.append((cx * cw, cy * ch,
                                              (cx + 1) * cw, (cy + 1) * ch))
            except Exception:  # noqa: BLE001
                pass
            if ratio >= 0.05:
                changed_pages.append(i + 1)

            canvas = Image.new("RGB", (w * 2 + GAP, h + BAND), (255, 255, 255))
            canvas.paste(ia, (0, BAND))
            canvas.paste(ib, (w + GAP, BAND))
            dr = ImageDraw.Draw(canvas)
            dr.line((w + GAP / 2, 0, w + GAP / 2, h + BAND), fill=(200, 200, 200), width=2)
            for (x0, y0, x1, y1) in boxes:
                dr.rectangle((x0, BAND + y0, x1, BAND + y1), outline=(255, 0, 0), width=2)
                dr.rectangle((x0 + w + GAP, BAND + y0, x1 + w + GAP, BAND + y1),
                             outline=(255, 0, 0), width=2)
            status = f"BEDA {ratio}%" if ratio >= 0.05 else "SAMA"
            color = (200, 0, 0) if ratio >= 0.05 else (0, 130, 0)
            dr.rectangle((0, 0, w * 2 + GAP, BAND), fill=(245, 245, 245))
            dr.text((8, 7), f"Hal {i + 1}  |  {status}", fill=color)

            page = out.new_page(width=canvas.width, height=canvas.height)
            page.insert_image(page.rect, stream=_png_bytes(canvas))
        out.save(output, garbage=4, deflate=True)
        out.close()
    return {"output": output, "pages": total, "changed": changed_pages}


def _png_bytes(im: Image.Image) -> bytes:
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


# ---------------------------------------------------------------- form fields
def list_form_fields(pdf_path: str) -> list[dict]:
    """Daftar field AcroForm: {name, type, value, page}."""
    _require_exists(pdf_path)
    fields: list[dict] = []
    with pymupdf.open(pdf_path) as doc:
        for pno, page in enumerate(doc):
            for w in page.widgets() or []:
                fields.append({
                    "name": w.field_name or "",
                    "type": w.field_type_string or "?",
                    "value": w.field_value,
                    "page": pno + 1,
                })
    return fields


def fill_form(pdf_path: str, output: str, values: dict) -> dict:
    """Isi field AcroForm dari dict {nama: nilai}. Return jumlah terisi."""
    _require_exists(pdf_path)
    _ensure_parent(output)
    filled = 0
    with pymupdf.open(pdf_path) as doc:
        if doc.is_encrypted:
            raise ConvertError("PDF terkunci — buka dulu lewat Unlock.")
        for page in doc:
            for w in page.widgets() or []:
                name = w.field_name or ""
                if name in (values or {}):
                    try:
                        w.field_value = str(values[name])
                        w.update()
                        filled += 1
                    except Exception:  # noqa: BLE001
                        continue
        doc.need_appearances = True
        doc.save(output, garbage=4, deflate=True)
    return {"output": output, "filled": filled}


# ---------------------------------------------------------------- scan -> PDF
def images_to_pdf_scan(images: list[str], output: str, gray: bool = True) -> str:
    """Gambar (hasil scan) -> PDF 1 halaman/gambar, opsi hitam-putih + kontras."""
    if not images:
        raise ConvertError("Belum ada gambar.")
    _ensure_parent(output)
    from PIL import ImageOps
    out = pymupdf.open()
    try:
        for p in images:
            _require_exists(p)
            im = Image.open(p)
            if gray:
                im = ImageOps.autocontrast(im.convert("L"))
                im = im.convert("RGB")
            else:
                im = ImageOps.autocontrast(im.convert("RGB"))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=85)
            w, h = im.size
            pw, ph = 595.0, 842.0  # A4 portrait
            s = min(pw / w, ph / h)
            iw, ih = w * s, h * s
            page = out.new_page(width=pw, height=ph)
            page.insert_image(
                pymupdf.Rect((pw - iw) / 2, (ph - ih) / 2,
                             (pw + iw) / 2, (ph + ih) / 2),
                stream=buf.getvalue())
        out.save(output, garbage=4, deflate=True)
    finally:
        out.close()
    return output
