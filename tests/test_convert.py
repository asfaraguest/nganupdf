"""Test konversi offline (headless, tanpa tkinter).

Uji: PDF->Word/PPT/Excel, PDF/A, ekstrak gambar, banding, form field, scan.
Dependensi opsional (docx/pptx/openpyxl) dilewati bila tidak terpasang,
tetapi di CI/dev seharusnya ada. Ghostscript diuji bila tersedia.
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import pymupdf
from PIL import Image

from core import convert_ops

TMP = os.path.join(os.path.dirname(__file__), "tmp")
os.makedirs(TMP, exist_ok=True)


def _has(mod):
    try:
        __import__(mod)
        return True
    except ImportError:
        return False


def make_pdf(path, pages=2, with_img=True, with_field=True):
    img = os.path.join(TMP, "cv.png")
    if not os.path.isfile(img):
        Image.new("RGB", (120, 90), (20, 120, 210)).save(img)
    d = pymupdf.open()
    for i in range(pages):
        pg = d.new_page()
        pg.insert_text((72, 72), f"Halo Halaman {i + 1}")
        pg.insert_text((72, 100), "Baris kedua untuk Word.")
    if with_img:
        d[0].insert_image(pymupdf.Rect(72, 140, 192, 230), filename=img)
    if with_field:
        w = pymupdf.Widget()
        w.field_name = "nama"
        w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
        w.rect = pymupdf.Rect(72, 260, 300, 282)
        w.field_value = ""
        d[0].add_widget(w)
    d.save(path)
    d.close()
    return path


def main():
    src = make_pdf(os.path.join(TMP, "conv-a.pdf"))

    if _has("docx"):
        out = os.path.join(TMP, "cv.docx")
        convert_ops.pdf_to_word(src, out)
        assert os.path.getsize(out) > 0
        from docx import Document
        text = "\n".join(p.text for p in Document(out).paragraphs)
        assert "Halo Halaman 1" in text, "teks Word hilang"
        print("pdf_to_word OK")
    else:
        print("pdf_to_word DILEWATI (python-docx absen)")

    if _has("pptx"):
        out = os.path.join(TMP, "cv.pptx")
        convert_ops.pdf_to_powerpoint(src, out, dpi=72)
        from pptx import Presentation
        prs = Presentation(out)
        assert len(list(prs.slides)) == 2, "jumlah slide salah"
        print("pdf_to_powerpoint OK")
    else:
        print("pdf_to_powerpoint DILEWATI (python-pptx absen)")

    if _has("openpyxl"):
        out = os.path.join(TMP, "cv.xlsx")
        r = convert_ops.pdf_to_excel(src, out)
        assert r["sheets"] >= 1 and os.path.getsize(out) > 0
        print(f"pdf_to_excel OK ({r['sheets']} sheet)")
    else:
        print("pdf_to_excel DILEWATI (openpyxl absen)")

    # PDF/A: ghostscript bila ada, else best-effort (tetap keluar file)
    out = os.path.join(TMP, "cv-pdfa.pdf")
    r = convert_ops.to_pdfa(src, out)
    assert os.path.isfile(out) and os.path.getsize(out) > 0
    print(f"to_pdfa OK method={r['method']} valid={r['valid']}")

    # ekstrak gambar
    imgs = convert_ops.extract_images(src, os.path.join(TMP, "cv-imgs"))
    assert len(imgs) >= 1 and all(os.path.isfile(p) for p in imgs)
    print(f"extract_images OK ({len(imgs)})")

    # banding: b beda + 1 halaman
    b = os.path.join(TMP, "conv-b.pdf")
    make_pdf(b, pages=1, with_img=False, with_field=False)
    out = os.path.join(TMP, "cv-cmp.pdf")
    r = convert_ops.compare_pdfs(src, b, out, dpi=72)
    assert os.path.isfile(out) and r["pages"] == 2
    assert 1 in r["changed"], "halaman 1 harus terdeteksi beda"
    print(f"compare_pdfs OK changed={r['changed']}")

    # form field
    fields = convert_ops.list_form_fields(src)
    assert any(f["name"] == "nama" for f in fields), "field tak terdeteksi"
    out = os.path.join(TMP, "cv-filled.pdf")
    r = convert_ops.fill_form(src, out, {"nama": "Budi"})
    assert r["filled"] == 1
    after = {f["name"]: f["value"] for f in convert_ops.list_form_fields(out)}
    assert after.get("nama") == "Budi", "nilai field tak tersimpan"
    print("form fields OK")

    # scan -> PDF
    img = os.path.join(TMP, "cv.png")
    out = os.path.join(TMP, "cv-scan.pdf")
    convert_ops.images_to_pdf_scan([img, img], out, gray=True)
    assert pymupdf.open(out).page_count == 2
    print("images_to_pdf_scan OK")

    print("\nTEST CONVERT LOLOS")


if __name__ == "__main__":
    main()
